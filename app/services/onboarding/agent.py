
"""
Onboarding Agent - Tenant SOP Analysis & Configuration Generator

Handles the full onboarding workflow:
  Task 1: Subquery refactoring based on SOP analysis
  Task 2: Agent 3 prompt customization (Summarization Agent)
  Task 3: Agent 4 prompt customization (Summary Validator Agent)

Stores all configurations in MongoDB tenant domain schema.
Supports re-onboarding when SOP is updated.
"""

import logging
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional
from pymongo import MongoClient
from app.core.config import settings
from app.services.summarization.prompts import (
    SUBQUERIES,
    SUMMARY_VALIDATOR_SYSTEM_PROMPT,
    MAIN_SUMMARY_INSTRUCTIONS,
    DEFAULT_SUMMARY_FORMAT,
    MAIN_SUMMARY_SYSTEM_PROMPT,
)
import openai
import json

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────
# Prompt Templates for the Onboarding Agent
# ─────────────────────────────────────────────

SUBQUERY_REFACTOR_SYSTEM_PROMPT = """
You are an expert AI systems analyst specializing in financial document analysis pipelines.

Your task is to analyze a tenant's Standard Operating Procedure (SOP) document and compare it against a set of DEFAULT SUBQUERIES used for retrieving information from DRHP/RHP documents stored in a vector database.

## DEFAULT SUBQUERIES (Base Template):
{default_subqueries}

## RULES:
1. The 10 default subqueries form the BASE template. They must NOT be removed unless the SOP explicitly excludes a topic.
2. Each default subquery can be REWORDED to better align with the SOP's terminology and focus areas.
3. NEW subqueries can be ADDED if the SOP demands coverage of topics not addressed by existing subqueries.
4. Identify any domain-specific terminology from the SOP that should be incorporated into the subqueries.
5. Ensure the subqueries cover ALL extraction areas mentioned in the SOP.

## OUTPUT FORMAT:
Return a JSON object with this exact structure:
{{
    "analysis": {{
        "missing_areas": ["List of topics in SOP not covered by default subqueries"],
        "additional_expectations": ["List of extra data points the SOP requires"],
        "domain_terminology": ["List of domain-specific terms found in SOP"]
    }},
    "subqueries": [
        "Subquery 1 text (modified or original)",
        "Subquery 2 text (modified or original)",
        ...
        "Any additional subqueries"
    ],
    "changes_log": [
        "Description of each change made and why"
    ]
}}

IMPORTANT: Return ONLY valid JSON. No markdown formatting, no code blocks.
"""

AGENT3_PROMPT_CUSTOMIZATION_SYSTEM_PROMPT = """
You are an expert AI prompt engineer specializing in financial document summarization systems.

Your task is to customize the SUMMARIZATION AGENT's system prompt based on a tenant's Standard Operating Procedure (SOP).

## BASE SUMMARIZATION PROMPT:
{base_prompt}

## CUSTOMIZATION RULES:
1. Preserve the core summarization capabilities and accuracy requirements.
2. Modify the section structure to match the SOP's required headings and hierarchy.
3. Adjust data ordering to match SOP requirements.
4. Include any mandatory disclosures specified by the SOP.
5. Adapt the tone and language to match the SOP's domain language.
6. Ensure all SOP-specified data points are explicitly requested in the prompt.
7. Maintain all table formats but adjust columns/rows per SOP requirements.
8. Keep the critical accuracy rules (zero fabrication, exact transcription, etc.)

## OUTPUT:
Return ONLY the complete, customized system prompt text. No explanations, no markdown code blocks.
The output must be a ready-to-use system prompt for the summarization agent.
"""

AGENT4_PROMPT_CUSTOMIZATION_SYSTEM_PROMPT = """
You are an expert AI prompt engineer specializing in financial document validation systems.

Your task is to customize the VALIDATION AGENT's system prompt based on a tenant's Standard Operating Procedure (SOP).

## BASE VALIDATION PROMPT:
{base_prompt}

## CUSTOMIZATION RULES:
1. Retain the core identity and strict validation rules (accuracy, verification, zero fabrication).
2. Replace generic format checks with SPECIFIC checks for the sections defined in the SOP.
3. Ensure the validator checks for EVERY section, table, and data point explicitly mentioned in the SOP.
4. Add SOP-driven validation changes such as:
   - Mandatory financial tables checks
   - Investor disclosure verification
   - Risk factor completeness checks
   - Regulatory formatting requirements
5. Maintain the systematic validation workflow (audit → cross-verification → compliance → missing data → reconstruction → QA).
6. Update the section checklists to match the customized section structure.

## OUTPUT:
Return ONLY the complete, customized system prompt text. No explanations, no markdown code blocks.
The output must be a ready-to-use system prompt for the validation agent.
"""


class OnboardingAgent:
    """
    Handles tenant onboarding by analyzing SOP documents and generating
    customized pipeline configurations (subqueries, Agent 3 prompt, Agent 4 prompt).
    
    Stores all configs in MongoDB under the tenant's domain document.
    Supports re-onboarding when SOP is updated.
    """

    def __init__(self):
        self.client = MongoClient(settings.MONGO_URI)
        self.db = self.client[settings.MONGO_DB_NAME]
        self.collection = self.db["domains"]
        self.openai_client = openai.OpenAI(api_key=settings.OPENAI_API_KEY)

    # ─────────────────────────────────────────────
    # File Extraction Utility
    # ─────────────────────────────────────────────

    def extract_text(self, file_content: bytes, filename: str) -> str:
        """Extracts text from PDF or DOCX file content."""
        import io
        text = ""
        try:
            if filename.lower().endswith(".pdf"):
                from PyPDF2 import PdfReader
                reader = PdfReader(io.BytesIO(file_content))
                for page in reader.pages:
                    text += page.extract_text() + "\n"
            elif filename.lower().endswith(".docx"):
                from docx import Document
                doc = Document(io.BytesIO(file_content))
                for para in doc.paragraphs:
                    text += para.text + "\n"
            else:
                # Assume plain text
                text = file_content.decode("utf-8", errors="ignore")
        except Exception as e:
            logger.error(f"Error extracting text from {filename}: {e}")
            return ""
        return text

    # ─────────────────────────────────────────────
    # Task 0: SOP Text Analysis (Raw → Structured)
    # ─────────────────────────────────────────────

    def analyze_and_create_format_prompt(self, raw_sop_text: str) -> str:
        """
        Analyzes RAW SOP text (extracted from file) and converts it into a
        clean, structured PROMPT TEMPLATE similar to DEFAULT_SUMMARY_FORMAT.
        """
        system_instructions = """
        You are an expert Prompt Engineer.
        Your task is to convert a raw "Standard Operating Procedure" (SOP) document into a specialized "REQUIRED FORMAT AND STRUCTURE" prompt template.

        GOAL:
        Create a clean, markdown-formatted template that an AI agent uses to structure a financial summary.
        The output must be ONLY the format/structure part (like headers, bullet points, tables with placeholders).
        
        RULES:
        1. Read the raw SOP content.
        2. Identify all required sections, tables, and data points.
        3. Convert them into a template format using Markdown.
        4. Use placeholders like [Amount], [Date], [%], etc.
        5. Do NOT include instructions on "how" to extract (unless critical formatting notes).
        6. Start with "## REQUIRED FORMAT AND STRUCTURE:"
        
        RAW SOP CONTENT:
        {raw_sop}

        OUTPUT:
        Return ONLY the generated template string.
        """
        try:
            response = self.openai_client.chat.completions.create(
                model="gpt-4.1-mini",
                messages=[
                    {"role": "system", "content": system_instructions.format(raw_sop=raw_sop_text)},
                    {"role": "user", "content": "Generate the format template."}
                ],
                temperature=0.2,
            )
            return response.choices[0].message.content.strip()
        except Exception as e:
            logger.error(f"Failed to generate format prompt: {e}")
            return ""

    # ─────────────────────────────────────────────
    # Task 1: Subquery Refactoring
    # ─────────────────────────────────────────────

    def _task1_refactor_subqueries(self, sop_text: str) -> Dict[str, Any]:
        """
        Task 1: Compare tenant SOP vs existing subqueries.
        Identify missing areas, modify/extend subqueries as needed.
        
        Returns:
            {
                "subqueries": [...],
                "analysis": {...},
                "changes_log": [...]
            }
        """
        logger.info("Task 1: Subquery Refactoring - Starting")

        # Format default subqueries for the prompt
        default_sq_formatted = "\n".join(
            [f"{i+1}. {sq}" for i, sq in enumerate(SUBQUERIES)]
        )

        try:
            response = self.openai_client.chat.completions.create(
                model="gpt-4.1-mini",
                messages=[
                    {
                        "role": "system",
                        "content": SUBQUERY_REFACTOR_SYSTEM_PROMPT.format(
                            default_subqueries=default_sq_formatted
                        ),
                    },
                    {
                        "role": "user",
                        "content": f"Analyze this tenant SOP and refactor subqueries:\n\n{sop_text}",
                    },
                ],
                temperature=0.2,
                response_format={"type": "json_object"},
            )

            result = json.loads(response.choices[0].message.content)
            
            # Validate: Must have at least the default number of subqueries
            subqueries = result.get("subqueries", [])
            if len(subqueries) < len(SUBQUERIES):
                logger.warning(
                    f"Task 1: Generated fewer subqueries ({len(subqueries)}) than default ({len(SUBQUERIES)}). "
                    "Padding with defaults."
                )
                # Pad with any missing default subqueries
                for i in range(len(subqueries), len(SUBQUERIES)):
                    subqueries.append(SUBQUERIES[i])
                result["subqueries"] = subqueries

            logger.info(
                f"Task 1: Completed. Generated {len(subqueries)} subqueries. "
                f"Changes: {len(result.get('changes_log', []))}"
            )
            return result

        except Exception as e:
            logger.error(f"Task 1: Failed - {e}. Using default subqueries.")
            return {
                "subqueries": list(SUBQUERIES),  # Copy of defaults
                "analysis": {"error": str(e)},
                "changes_log": ["Fallback: Using default subqueries due to error"],
            }

    # ─────────────────────────────────────────────
    # Task 2: Agent 3 Prompt Customization
    # ─────────────────────────────────────────────

    def _task2_customize_agent3_prompt(self, sop_text: str) -> str:
        """
        Task 2: Customize the summarization agent (Agent 3) prompt
        based on the tenant's SOP requirements.
        
        Returns: Customized Agent 3 system prompt string.
        """
        logger.info("Task 2: Agent 3 Prompt Customization - Starting")

        try:
            response = self.openai_client.chat.completions.create(
                model="gpt-4.1-mini",
                messages=[
                    {
                        "role": "system",
                        "content": AGENT3_PROMPT_CUSTOMIZATION_SYSTEM_PROMPT.format(
                            base_prompt=MAIN_SUMMARY_SYSTEM_PROMPT
                        ),
                    },
                    {
                        "role": "user",
                        "content": f"Customize the summarization agent prompt based on this SOP:\n\n{sop_text}",
                    },
                ],
                temperature=0.2,
            )

            custom_prompt = response.choices[0].message.content.strip()
            logger.info(
                f"Task 2: Completed. Custom Agent 3 prompt length: {len(custom_prompt)} chars"
            )
            return custom_prompt

        except Exception as e:
            logger.error(f"Task 2: Failed - {e}. Using default Agent 3 prompt.")
            return MAIN_SUMMARY_SYSTEM_PROMPT

    # ─────────────────────────────────────────────
    # Task 3: Agent 4 Prompt Customization
    # ─────────────────────────────────────────────

    def _task3_customize_agent4_prompt(self, sop_text: str) -> str:
        """
        Task 3: Customize the validator agent (Agent 4) prompt
        based on the tenant's SOP validation rules.
        
        Returns: Customized Agent 4 system prompt string.
        """
        logger.info("Task 3: Agent 4 Prompt Customization - Starting")

        try:
            response = self.openai_client.chat.completions.create(
                model="gpt-4.1-mini",
                messages=[
                    {
                        "role": "system",
                        "content": AGENT4_PROMPT_CUSTOMIZATION_SYSTEM_PROMPT.format(
                            base_prompt=SUMMARY_VALIDATOR_SYSTEM_PROMPT
                        ),
                    },
                    {
                        "role": "user",
                        "content": f"Customize the validator agent prompt based on this SOP:\n\n{sop_text}",
                    },
                ],
                temperature=0.2,
            )

            custom_prompt = response.choices[0].message.content.strip()
            logger.info(
                f"Task 3: Completed. Custom Agent 4 prompt length: {len(custom_prompt)} chars"
            )
            return custom_prompt

        except Exception as e:
            logger.error(f"Task 3: Failed - {e}. Using default Agent 4 prompt.")
            return SUMMARY_VALIDATOR_SYSTEM_PROMPT

    # ─────────────────────────────────────────────
    # Main Onboarding Orchestrator
    # ─────────────────────────────────────────────

    def process_new_tenant(
        self,
        domain_id: str,
        custom_sop_input: str,
        is_raw_text: bool = False,
        toggles: dict = None,
    ) -> bool:
        """
        Processes a new tenant's onboarding configuration.
        Runs all 3 tasks sequentially, keeping SOP context throughout.
        Supports both initial onboarding and re-onboarding.

        Args:
            domain_id: Tenant domain ID
            custom_sop_input: Either the final template or RAW SOP text/file content
            is_raw_text: True if input is raw text that needs AI analysis first
            toggles: Feature flags (investor_match_only, valuation_matching, etc.)
        
        Returns:
            True if successfully onboarded, False otherwise.
        """
        if toggles is None:
            toggles = {}

        logger.info(f"═══════════════════════════════════════════════")
        logger.info(f"Onboarding Agent: Starting for domain {domain_id}")
        logger.info(f"═══════════════════════════════════════════════")

        # ── Step 0: If raw text (from file upload), convert to structured SOP ──
        final_sop_text = custom_sop_input
        if is_raw_text and custom_sop_input:
            logger.info("Step 0: Converting raw SOP text to structured template...")
            generated_format = self.analyze_and_create_format_prompt(custom_sop_input)
            if generated_format:
                final_sop_text = generated_format
                logger.info(f"Step 0: Generated structured SOP ({len(final_sop_text)} chars)")
            else:
                logger.warning("Step 0: Failed to generate structured SOP, using raw text")

        # If no SOP provided, store toggles only with empty SOP fields
        if not final_sop_text or not final_sop_text.strip():
            logger.info("No SOP provided. Storing toggles only, using default pipeline.")
            update_data = {
                "sop_text": "",
                "custom_summary_sop": "",
                "custom_subqueries": [],
                "agent3_prompt": "",
                "agent4_prompt": "",
                "custom_validator_prompt": "",
                "onboarding_status": "completed_no_sop",
                "last_onboarded": datetime.now(timezone.utc).isoformat(),
                "investor_match_only": toggles.get("investor_match_only", False),
                "valuation_matching": toggles.get("valuation_matching", False),
                "adverse_finding": toggles.get("adverse_finding", False),
            }
            if "target_investors" in toggles:
                update_data["target_investors"] = toggles["target_investors"]

            return self._save_to_mongodb(domain_id, update_data)

        # ── Task 1: Subquery Refactoring ──
        logger.info("─── Task 1/3: Subquery Refactoring ───")
        subquery_result = self._task1_refactor_subqueries(final_sop_text)
        custom_subqueries = subquery_result.get("subqueries", list(SUBQUERIES))

        # ── Task 2: Agent 3 Prompt Customization ──
        logger.info("─── Task 2/3: Agent 3 Prompt Customization ───")
        agent3_prompt = self._task2_customize_agent3_prompt(final_sop_text)

        # ── Task 3: Agent 4 Prompt Customization ──
        logger.info("─── Task 3/3: Agent 4 Prompt Customization ───")
        agent4_prompt = self._task3_customize_agent4_prompt(final_sop_text)

        # ── Store Everything in MongoDB ──
        logger.info("Storing onboarding configuration in MongoDB...")
        update_data = {
            # SOP Storage
            "sop_text": custom_sop_input,        # Original uploaded SOP text
            "custom_summary_sop": final_sop_text, # Structured/processed SOP

            # Task 1 output
            "custom_subqueries": custom_subqueries,
            "subquery_analysis": subquery_result.get("analysis", {}),
            "subquery_changes_log": subquery_result.get("changes_log", []),

            # Task 2 output
            "agent3_prompt": agent3_prompt,

            # Task 3 output  
            "agent4_prompt": agent4_prompt,
            "custom_validator_prompt": agent4_prompt,  # backward compat field

            # Toggles
            "investor_match_only": toggles.get("investor_match_only", False),
            "valuation_matching": toggles.get("valuation_matching", False),
            "adverse_finding": toggles.get("adverse_finding", False),

            # Metadata
            "onboarding_status": "completed",
            "last_onboarded": datetime.now(timezone.utc).isoformat(),
        }

        if "target_investors" in toggles:
            update_data["target_investors"] = toggles["target_investors"]

        success = self._save_to_mongodb(domain_id, update_data)

        if success:
            logger.info(f"═══════════════════════════════════════════════")
            logger.info(f"Onboarding Agent: COMPLETED for domain {domain_id}")
            logger.info(f"  Subqueries: {len(custom_subqueries)} (default: {len(SUBQUERIES)})")
            logger.info(f"  Agent 3 Prompt: {len(agent3_prompt)} chars")
            logger.info(f"  Agent 4 Prompt: {len(agent4_prompt)} chars")
            logger.info(f"═══════════════════════════════════════════════")
        else:
            logger.error(f"Onboarding Agent: FAILED for domain {domain_id}")

        return success

    # ─────────────────────────────────────────────
    # MongoDB Storage
    # ─────────────────────────────────────────────

    def _save_to_mongodb(self, domain_id: str, update_data: Dict[str, Any]) -> bool:
        """Saves onboarding configuration to MongoDB (upsert)."""
        try:
            result = self.collection.update_one(
                {"domainId": domain_id},
                {"$set": update_data},
                upsert=True,
            )
            logger.info(
                f"MongoDB update for {domain_id}: "
                f"matched={result.matched_count}, modified={result.modified_count}, "
                f"upserted_id={result.upserted_id}"
            )
            return True
        except Exception as e:
            logger.error(f"Failed to save onboarding config to MongoDB: {e}")
            return False


# ─────────────────────────────────────────────
# Public API Functions
# ─────────────────────────────────────────────

def onboard_tenant(
    domain_id: str,
    custom_sop_input: str,
    is_raw_text: bool,
    toggles: dict,
) -> bool:
    """
    Public function to trigger tenant onboarding.
    Called from API endpoint or scripts.
    """
    agent = OnboardingAgent()
    return agent.process_new_tenant(domain_id, custom_sop_input, is_raw_text, toggles)
